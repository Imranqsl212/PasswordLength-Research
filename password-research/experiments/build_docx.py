"""Build the portable, editable Word manuscript from canonical Markdown.

No external document account or plaintext credential input is used. Figure
bytes are embedded in OOXML; short equations are editable Unicode paragraphs.
"""
from pathlib import Path
from datetime import datetime, timezone
import argparse
import hashlib
import json
import re
import zipfile

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "paper/submission/Imran_Zakirov_Password_Length_Manuscript.docx"


def hyperlink(paragraph, label, url, bold=False):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle"); style.set(qn("w:val"), "Hyperlink"); props.append(style)
    if bold:
        b = OxmlElement("w:b"); props.append(b)
    run.append(props)
    text = OxmlElement("w:t"); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def inline(paragraph, value):
    pattern = r"(\[[^\]]+\]\([^\n)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)"
    for token in re.split(pattern, value):
        if not token:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
        if match:
            label, target = match.groups()
            if re.match(r"https?://", target):
                if label.isdigit():
                    label = "[" + label + "]"
                hyperlink(paragraph, label, target)
            else:
                paragraph.add_run(label)
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def add_equation(document, value):
    p = document.add_paragraph(style="Equation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # These short Unicode equations remain editable and visible in readers
    # whose previews omit Office Math objects. They need no linked fonts/images.
    r = p.add_run(value)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11.5)
    return p


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    props = cell._tc.get_or_add_tcPr()
    tcw = props.find(qn("w:tcW"))
    tcw.set(qn("w:w"), str(round(inches * 1440))); tcw.set(qn("w:type"), "dxa")


def add_table(document, rows):
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r":?-+:?", c or " ") for c in row)]
    n = len(cells[0])
    if any(len(row) != n for row in cells):
        raise ValueError("Malformed Markdown table")
    table = document.add_table(rows=0, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    table.style = "Light Shading Accent 1"
    # First explanatory column often needs more space than numerical columns.
    first = 2.15 if n >= 5 else 2.05 if n == 3 else 1.6
    widths = [first] + [(6.55-first)/(n-1)]*(n-1) if n > 1 else [6.55]
    for index, values in enumerate(cells):
        row = table.add_row()
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit"); props.append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader"); props.append(repeat)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_width(cell, widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.line_spacing = 1.0
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9 if n >= 6 else 9.5)
                if index == 0:
                    run.bold = True
            if j and re.fullmatch(r"[\d.,+−–% /×⁻⁰¹²³⁴⁵⁶⁷⁸⁹e]+", re.sub(r"\*", "", value)):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText"); code.set(qn("xml:space"), "preserve"); code.text = instruction
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, code, end])


def build(source=ROOT / "paper/manuscript.md", output=OUTPUT):
    document = Document()
    sec = document.sections[0]
    sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
    sec.top_margin = Inches(.78); sec.bottom_margin = Inches(.78)
    sec.left_margin = Inches(.86); sec.right_margin = Inches(.86)
    sec.header_distance = Inches(.3); sec.footer_distance = Inches(.3)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"; style.font.size = Pt(14 if level == 1 else 12)
        style.font.color.rgb = RGBColor.from_string("183E50"); style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5); style.paragraph_format.keep_with_next = True
    for name in ("Caption", "Reference", "Equation", "Code"):
        if name not in document.styles:
            document.styles.add_style(name, 1)
        style = document.styles[name]; style.font.name = "Times New Roman"
        style.font.size = Pt(9.5 if name == "Caption" else 10.5 if name == "Reference" else 11)
        style.paragraph_format.space_after = Pt(6)
    document.styles["Caption"].paragraph_format.keep_together = True
    document.styles["Reference"].paragraph_format.line_spacing = 1
    document.styles["Reference"].paragraph_format.left_indent = Inches(.27)
    document.styles["Reference"].paragraph_format.first_line_indent = Inches(-.27)
    document.styles["Code"].font.name = "Consolas"; document.styles["Code"].font.size = Pt(9)
    document.styles["Code"].paragraph_format.line_spacing = 1
    hyperlink_style = document.styles["Hyperlink"] if "Hyperlink" in document.styles else document.styles.add_style("Hyperlink", 2)
    hyperlink_style.font.color.rgb = RGBColor.from_string("1B6685")
    header = sec.header.paragraphs[0]
    header.add_run("Zakirov | Password length and frequency concentration").font.size = Pt(8)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ").font.size = Pt(8); field(footer, "PAGE")
    props = document.core_properties
    props.author = "Imran Zakirov"
    props.title = "Password Length and Frequency Concentration in Differentially Private Releases"
    props.subject = "Reproducible reanalysis with controlled generation baselines"
    props.keywords = "password length, frequency concentration, differential privacy, reproducibility"
    props.comments = "AI-assisted research; transparent disclosure is included in the manuscript."
    props.created = datetime.now(timezone.utc)
    props.modified = props.created

    lines = source.read_text().splitlines()
    images = []; tables = 0; references = 0; equations = 0; in_refs = False; in_code = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("```"):
            in_code = not in_code; i += 1; continue
        if in_code:
            p = document.add_paragraph(line, style="Code"); i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            add_table(document, rows); tables += 1; continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt, target = image_match.groups(); path = Path(target)
            if not path.is_absolute():
                path = source.parent / path
            if not path.exists():
                raise ValueError(f"Missing figure {path}")
            p = document.add_paragraph(); p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(); run.add_picture(str(path), width=Inches(6.55))
            extent = run._r.xpath(".//wp:docPr")[0]; extent.set("descr", alt)
            images.append({"file": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "alt": alt})
            i += 1; continue
        if line.startswith("$$") and line.endswith("$$"):
            add_equation(document, line[2:-2]); equations += 1; i += 1; continue
        heading = re.match(r"^(#{1,4}) (.+)$", line)
        if heading:
            depth, value = len(heading.group(1)), heading.group(2)
            if depth == 1:
                p = document.add_paragraph()
                r = p.add_run(value); r.bold = True; r.font.size = Pt(20)
                p.paragraph_format.space_after = Pt(8); p.paragraph_format.keep_with_next = True
            else:
                p = document.add_paragraph(value, style=f"Heading {min(depth-1,3)}")
            in_refs = value == "References"
            i += 1; continue
        if line.startswith("**Figure "):
            p = document.add_paragraph(style="Caption"); inline(p, line)
        elif line.startswith("**Table "):
            p = document.add_paragraph(style="Caption"); p.paragraph_format.keep_with_next = True; inline(p, line)
        elif in_refs and re.match(r"^\[\d+\] ", line):
            p = document.add_paragraph(style="Reference"); inline(p, line); references += 1
        elif line.startswith("- "):
            p = document.add_paragraph(style="List Bullet"); inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = document.add_paragraph(style="List Number"); inline(p, re.sub(r"^\d+\. ", "", line))
        else:
            p = document.add_paragraph(); inline(p, line)
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    with zipfile.ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        xml = archive.read("word/document.xml").decode()
        rels = archive.read("word/_rels/document.xml.rels").decode()
        if 'TargetMode="External"' in rels and re.search(r'Target="(?:file:|/Users/)', rels):
            raise ValueError("Document has nonportable external local relationships")
        if len(media) != len(images):
            raise ValueError("Every figure must have distinct embedded media bytes")
        if "DATA UNAVAILABLE" not in xml or "NOT RUN" not in xml:
            raise ValueError("Missing explicit experiment-status text")
        if "UVK" not in xml or "Imran Zakirov" not in xml:
            raise ValueError("Author metadata/text missing")
    audit = {"docx": str(output.relative_to(ROOT)), "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
             "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "python_docx": docx.__version__,
             "builder_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
             "figures_embedded": len(images), "tables": tables, "references": references,
             "editable_equation_paragraphs": equations, "figures": images,
             "local_external_relationships": 0,
             "layout_render": "UNVERIFIED: actual Word/Pages pagination; structural audit only at build time"}
    (ROOT / "results/statistics/docx_audit.json").write_text(json.dumps(audit, indent=2)+"\n")
    print(json.dumps({k:v for k,v in audit.items() if k != "figures"}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--source", type=Path); parser.add_argument("--output", type=Path)
    args = parser.parse_args(); build(args.source or ROOT / "paper/manuscript.md", args.output or OUTPUT)
